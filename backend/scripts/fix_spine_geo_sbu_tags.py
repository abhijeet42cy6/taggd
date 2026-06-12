#!/usr/bin/env python3
"""
Clear hierarchy_tag_sbu in spine workbooks when it duplicates sub_region (geographic data).

Preserves row 1 headers and row 2 hints; edits data from row 3 onward.
Also clears paired client hierarchy_tag_sbu when it matches the project sub_region on the same row.

Run from repo root:
  python3 backend/scripts/fix_spine_geo_sbu_tags.py excel_files_imp/01_spine_clients_projects_filled.xlsx
  python3 backend/scripts/fix_spine_geo_sbu_tags.py path/to/workbook.xlsx --dry-run
"""
from __future__ import annotations

import argparse
import os
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "../..")))


def _col_index(header_row, name: str) -> int | None:
    target = name.strip().lower()
    for idx, cell in enumerate(header_row, start=1):
        val = cell.value if hasattr(cell, "value") else cell
        if val is not None and str(val).strip().lower() == target:
            return idx
    return None


def _cell_str(val) -> str:
    if val is None:
        return ""
    return str(val).strip()


def fix_workbook(path: str, *, dry_run: bool = False) -> dict[str, int]:
    wb = load_workbook(path)
    stats = {"projects_cleared": 0, "clients_cleared": 0}

    if "projects" not in wb.sheetnames or "clients" not in wb.sheetnames:
        raise SystemExit("Workbook must contain 'clients' and 'projects' sheets")

    ws_p = wb["projects"]
    ws_c = wb["clients"]
    p_header = [ws_p.cell(row=1, column=c).value for c in range(1, ws_p.max_column + 1)]
    c_header = [ws_c.cell(row=1, column=c).value for c in range(1, ws_c.max_column + 1)]

    p_sbu_col = _col_index(p_header, "hierarchy_tag_sbu")
    p_sub_col = _col_index(p_header, "sub_region")
    c_sbu_col = _col_index(c_header, "hierarchy_tag_sbu")

    if p_sbu_col is None or p_sub_col is None:
        raise SystemExit("projects sheet missing hierarchy_tag_sbu or sub_region column")
    if c_sbu_col is None:
        raise SystemExit("clients sheet missing hierarchy_tag_sbu column")

    max_row = max(ws_p.max_row, ws_c.max_row)
    for row in range(3, max_row + 1):
        sub = _cell_str(ws_p.cell(row=row, column=p_sub_col).value)
        p_sbu = _cell_str(ws_p.cell(row=row, column=p_sbu_col).value)
        c_sbu = _cell_str(ws_c.cell(row=row, column=c_sbu_col).value)

        if sub and p_sbu and p_sbu.lower() == sub.lower():
            stats["projects_cleared"] += 1
            if not dry_run:
                ws_p.cell(row=row, column=p_sbu_col).value = None

        if sub and c_sbu and c_sbu.lower() == sub.lower():
            stats["clients_cleared"] += 1
            if not dry_run:
                ws_c.cell(row=row, column=c_sbu_col).value = None

    if not dry_run:
        wb.save(path)

    return stats


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("workbook", help="Path to 01_spine_clients_projects*.xlsx")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    path = os.path.abspath(args.workbook)
    if not os.path.isfile(path):
        raise SystemExit(f"File not found: {path}")
    stats = fix_workbook(path, dry_run=args.dry_run)
    mode = "Would clear" if args.dry_run else "Cleared"
    print(f"{mode} hierarchy_tag_sbu on {stats['projects_cleared']} project rows")
    print(f"{mode} hierarchy_tag_sbu on {stats['clients_cleared']} client rows")
    if not args.dry_run:
        print(f"Saved {path}")


if __name__ == "__main__":
    main()
