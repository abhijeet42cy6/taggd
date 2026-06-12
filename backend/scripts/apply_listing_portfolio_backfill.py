#!/usr/bin/env python3
"""
Backfill projects.category / practice / vertical from excel_files_imp/Copy of Listing.xlsx.

Maps Listing columns → projects:
  Sub Category  → category  (TARA / Non TARA)
  Account Type  → practice  (Lateral / Leadership)
  Industry      → vertical

Does NOT change region, regional_head, practice_head, or account_status.

Usage (repo root):
  python3 backend/scripts/apply_listing_portfolio_backfill.py
  DATABASE_URL=postgresql+psycopg://... python3 backend/scripts/apply_listing_portfolio_backfill.py --apply
  python3 backend/scripts/apply_listing_portfolio_backfill.py --book excel_files_imp/Copy\\ of\\ Listing.xlsx --apply
"""
from __future__ import annotations

import argparse
import os
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

_REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(_REPO))

from backend.db.database import Project  # noqa: E402


def _norm_name(s: Optional[str]) -> str:
    if not s:
        return ""
    return re.sub(r"\s+", " ", str(s).strip().lower())


def _cell_str(v: Any) -> Optional[str]:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def load_listing(path: Path) -> List[Dict[str, Any]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise SystemExit(f"Empty workbook: {path}")
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    out: List[Dict[str, Any]] = []
    for raw in rows[1:]:
        if not raw or raw[0] is None or str(raw[0]).strip() == "":
            continue
        rec = {header[i]: raw[i] for i in range(min(len(header), len(raw)))}
        project = str(rec.get("Project") or "").strip()
        if not project:
            continue
        out.append(
            {
                "project": project,
                "sub_category": _cell_str(rec.get("Sub Category")),
                "account_type": _cell_str(rec.get("Account Type")),
                "industry": _cell_str(rec.get("Industry")),
            }
        )
    return out


def resolve_project_id(by_exact: Dict[str, int], by_norm: Dict[str, List[int]], name: str) -> Optional[int]:
    if name in by_exact:
        return by_exact[name]
    hits = by_norm.get(_norm_name(name), [])
    if len(hits) == 1:
        return hits[0]
    return None


def build_indexes(session: Session) -> Tuple[Dict[str, int], Dict[str, List[int]]]:
    by_exact: Dict[str, int] = {}
    by_norm: Dict[str, List[int]] = defaultdict(list)
    for p in session.query(Project).all():
        name = (p.account_name or "").strip()
        if name:
            by_exact[name] = p.id
            by_norm[_norm_name(name)].append(p.id)
    return by_exact, by_norm


def run_backfill(session: Session, listing_rows: List[Dict[str, Any]], apply: bool) -> Dict[str, Any]:
    by_exact, by_norm = build_indexes(session)
    missing: List[str] = []
    updated: List[Dict[str, Any]] = []
    unchanged = 0

    field_map = (
        ("sub_category", "category"),
        ("account_type", "practice"),
        ("industry", "vertical"),
    )

    for row in listing_rows:
        pid = resolve_project_id(by_exact, by_norm, row["project"])
        if pid is None:
            missing.append(row["project"])
            continue
        p = session.get(Project, pid)
        if not p:
            continue

        changes: Dict[str, Tuple[Any, Any]] = {}
        for src_key, dst_attr in field_map:
            new_val = row[src_key]
            old_val = getattr(p, dst_attr)
            old_s = (old_val or "").strip() if old_val is not None else None
            new_s = new_val
            if old_s == new_s:
                continue
            setattr(p, dst_attr, new_val)
            changes[dst_attr] = (old_val, new_val)

        if changes:
            updated.append(
                {
                    "project_id": pid,
                    "account_name": p.account_name,
                    "changes": {k: {"from": v[0], "to": v[1]} for k, v in changes.items()},
                }
            )
        else:
            unchanged += 1

    summary = {
        "listing_rows": len(listing_rows),
        "matched": len(listing_rows) - len(missing),
        "missing": missing,
        "updated": len(updated),
        "unchanged": unchanged,
        "updates": updated,
    }

    if apply:
        session.commit()
        print("==> COMMITTED to database")
    else:
        session.rollback()
        print("==> DRY RUN (no changes committed). Pass --apply to write.")

    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description="Backfill category/practice/vertical from Listing.xlsx")
    parser.add_argument(
        "--book",
        default=str(_REPO / "excel_files_imp" / "Copy of Listing.xlsx"),
        help="Path to Copy of Listing workbook",
    )
    parser.add_argument("--apply", action="store_true", help="Commit changes (default: dry-run)")
    args = parser.parse_args()

    book_path = Path(args.book)
    if not book_path.is_file():
        raise SystemExit(f"Workbook not found: {book_path}")

    db_url = os.environ.get("DATABASE_URL", "").strip()
    if not db_url:
        raise SystemExit("Set DATABASE_URL (postgresql+psycopg://...)")

    listing_rows = load_listing(book_path)
    print(f"==> Loaded {len(listing_rows)} rows from {book_path}")

    engine = create_engine(db_url, pool_pre_ping=True)
    SessionLocal = sessionmaker(bind=engine)
    session = SessionLocal()
    try:
        summary = run_backfill(session, listing_rows, apply=args.apply)
    finally:
        session.close()

    print("\n==> Summary")
    for k in ("listing_rows", "matched", "updated", "unchanged"):
        print(f"  {k}: {summary[k]}")
    if summary["missing"]:
        print(f"  missing: {summary['missing']}")
    if summary["updates"]:
        print("\n==> Sample updates (first 8)")
        for u in summary["updates"][:8]:
            print(f"  {u['account_name']} (id={u['project_id']}): {u['changes']}")

    if summary.get("missing"):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
