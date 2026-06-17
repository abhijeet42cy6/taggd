#!/usr/bin/env python3
"""
Build Maruti Suzuki Tracker - Production.xlsx from Maruti.xlsx (raw PM export).

Uses project_1_tracker_template (1).xlsx only as the Excel shell (headers + Contractual).

Usage (from revagent/):
  PYTHONPATH=. python3 excel_upload_masters/generate_maruti_production_tracker.py
"""
from __future__ import annotations

import datetime as dt
import re
import shutil
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
TAGGD_ROOT = ROOT.parent
SHELL = TAGGD_ROOT / "project_1_tracker_template (1).xlsx"
RAW_PM = TAGGD_ROOT / "Maruti.xlsx"
TEMPLATE_FALLBACK = TAGGD_ROOT / "Maruti Suzuki Tracker (3).xlsx"
OUT_NAME = "Maruti Suzuki Tracker - Production.xlsx"
OUT_PATH = TAGGD_ROOT / OUT_NAME
OUT_COPY = ROOT / "excel_files" / OUT_NAME

DATA_START = 4

_STATUS_PRIORITY = {
    "joined": 6,
    "ready to join": 5,
    "documentation": 4,
    "salary fitment": 4,
    "loi pending": 4,
    "interview": 3,
    "cv sent": 3,
    "sourcing": 2,
    "intake": 2,
    "final interview": 3,
    "on_hold": 1,
    "cancelled": 0,
}

_RAW_STATUS_MAP = {
    "joined": "Joined",
    "ready to join": "Offered",
    "loi pending": "Offered",
    "salary fitment": "Offered",
    "documentation": "Offered",
    "sourcing": "Open",
    "intake": "Open",
    "interview": "Open",
    "cv sent": "Open",
    "final interview": "Open",
    "on_hold": "On Hold",
    "cancelled": "Cancelled",
}

_RAW_STAGE_MAP = {
    "joined": "Joined",
    "offered": "Offered",
    "on_hold": "On Hold",
    "cancelled": "Cancelled",
    "in_process": "Open",
    "selection": "Offered",
}

HIRE_TYPE_COL = "Type of Hire\n(Lateral, AIT, Campus, IJP)"


def _is_empty(val) -> bool:
    if val is None:
        return True
    if isinstance(val, float) and pd.isna(val):
        return True
    s = str(val).strip()
    return s in ("", "nan", "None", "NaT", "-", "—")


def _norm_id(val) -> str:
    return str(val).strip().lstrip("\xa0")


def _norm_cand(val) -> str:
    s = str(val or "").strip().lower()
    return "" if s in ("", "-", "nan", "none") else s


def _to_date(val):
    if _is_empty(val):
        return None
    if isinstance(val, (dt.datetime, dt.date, pd.Timestamp)):
        return val.date() if hasattr(val, "date") else val
    s = str(val).strip()
    if s.replace(".", "", 1).isdigit():
        n = float(s)
        if n > 40000:
            return dt.date(1899, 12, 30) + dt.timedelta(days=int(n))
    try:
        parsed = pd.to_datetime(val, errors="coerce")
        if pd.notna(parsed):
            return parsed.date()
    except Exception:
        pass
    return None


def _ctc_lakhs(raw_ctc) -> float | None:
    if _is_empty(raw_ctc):
        return None
    try:
        v = float(str(raw_ctc).replace(",", ""))
    except ValueError:
        return None
    if v <= 0:
        return None
    if v > 500:
        return round(v / 100_000, 2)
    return round(v, 2)


def _map_source_joiner_type(source_of_hire: str, hire_type: str) -> str:
    soh = (source_of_hire or "").strip().lower()
    ht = (hire_type or "").strip().lower()
    if ht == "ijp" or soh == "ijp":
        return "IJP – Internal Job Posting"
    if soh in ("er", "employee referral"):
        return "ER – Employee Referral"
    if ht == "campus" or soh == "campus":
        return "Campus"
    if soh in ("rehire",) or "transfer" in ht:
        return "Internal Transfer"
    return "Taggd RPO"


def _map_status(raw_status: str, stage: str = "") -> str:
    st = (stage or "").strip().lower().replace(" ", "_")
    if st in _RAW_STAGE_MAP:
        return _RAW_STAGE_MAP[st]
    rs = (raw_status or "").strip().lower().replace(" ", "_")
    return _RAW_STATUS_MAP.get(rs, "Open")


def _status_priority(raw: pd.Series) -> int:
    rs = str(raw.get("Status") or "").strip().lower()
    st = str(raw.get("Stage") or "").strip().lower().replace(" ", "_")
    return max(_STATUS_PRIORITY.get(rs, 0), _STATUS_PRIORITY.get(st, 0))


def _header_map(ws) -> dict[str, int]:
    out: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        h = ws.cell(2, col).value
        if h and str(h) != "Scenario (example rows only)":
            out[str(h)] = col
    return out


def _clear_data_area(ws, from_row: int, to_row: int) -> None:
    for row in range(from_row, to_row + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row, col)
            if type(cell).__name__ != "MergedCell":
                cell.value = None


def _write_contractual(wb) -> None:
    ws = wb["Contractual"]
    rows = [
        (6, "All joiners (CTC below 40 Lakhs)", "1.10%", "2.40%", 1200),
        (7, "CTC 40 Lakhs and above", "1.00%", "2.20%", 1200),
    ]
    for r, band, open_pct, close_pct, flat in rows:
        for c, val in enumerate([band, open_pct, close_pct, flat], start=2):
            cell = ws.cell(r, c)
            if type(cell).__name__ != "MergedCell":
                cell.value = val


def _dedupe_raw(raw_df: pd.DataFrame) -> pd.DataFrame:
    """Keep best row per (Job ID, Candidate Name) — drops duplicate On_Hold echoes."""
    groups: dict[tuple[str, str], int] = {}
    for idx, raw in raw_df.iterrows():
        jid = _norm_id(raw.get("Job ID"))
        if not jid:
            continue
        cand = _norm_cand(raw.get("Candidate Name"))
        key = (jid, cand)
        if key not in groups or _status_priority(raw) > _status_priority(raw_df.loc[groups[key]]):
            groups[key] = idx
    kept = sorted(groups.values())
    return raw_df.loc[kept].reset_index(drop=True)


def _assign_req_ids(deduped: pd.DataFrame) -> list[str]:
    """One Req ID per row; suffix when multiple candidates share a Job ID."""
    job_counts: dict[str, int] = defaultdict(int)
    req_ids: list[str] = []
    for _, raw in deduped.iterrows():
        jid = _norm_id(raw.get("Job ID"))
        cand = _norm_cand(raw.get("Candidate Name"))
        job_counts[jid] += 1
    job_seen: dict[str, int] = defaultdict(int)
    for _, raw in deduped.iterrows():
        jid = _norm_id(raw.get("Job ID"))
        cand = _norm_cand(raw.get("Candidate Name"))
        if job_counts[jid] == 1:
            req_ids.append(jid)
        else:
            job_seen[jid] += 1
            req_ids.append(f"{jid}-{job_seen[jid]}")
    return req_ids


def _raw_row_to_template(raw: pd.Series, req_id: str) -> dict:
    status = _map_status(str(raw.get("Status") or ""), str(raw.get("Stage") or ""))
    ctc = _ctc_lakhs(raw.get("Offered CTC"))
    join_date = _to_date(raw.get("Joining Date"))
    selection = _to_date(raw.get("Selection Date"))
    loi = _to_date(raw.get("LOI Date"))
    req_created = _to_date(raw.get("Requisition Date")) or _to_date(raw.get("Job Creation Date"))
    mandate = _to_date(raw.get("Assigned Date to Taggd")) or req_created
    cand_raw = str(raw.get("Candidate Name") or "").strip()
    cand = cand_raw if not _is_empty(cand_raw) and cand_raw != "-" else ""
    if not cand and status == "Open":
        cand = f"Open — {req_id}"

    sjt = _map_source_joiner_type(
        str(raw.get("Source of Hire") or ""),
        str(raw.get(HIRE_TYPE_COL) or ""),
    )

    if status == "Joined" and ctc is None:
        ctc = 12.0  # fallback so revenue logic doesn't zero out

    return {
        "Req ID": req_id,
        "Position Title": str(raw.get("Job Name") or "").strip(),
        "Current Status": status,
        "Candidate Name": cand,
        "Hiring Manager": str(raw.get("Hiring Manager Name") or "").strip(),
        "Department": str(raw.get("Department") or "").strip(),
        "Location": str(raw.get("Job Location") or "").strip(),
        "Business Unit": str(raw.get("Business HRBP") or "").strip(),
        "Division": str(raw.get("Division ") or raw.get("Division") or "").strip(),
        "Job Type": str(raw.get("New / Replacement") or "").strip(),
        "Req Created Date": req_created,
        "Mandate Received Date": mandate,
        "Intake Date": _to_date(raw.get("Intake Date")),
        "Candidate Selection Date": selection,
        "Offered Date": loi or selection,
        "Offer Accepted Date": loi,
        "Joining Date": join_date,
        "LOI Date": loi,
        "Closure Date": join_date if status == "Joined" else None,
        "Offered CTC (Lakhs)": ctc,
        "Source Joiner Type": sjt,
        "Source of Hire": str(raw.get("Source of Hire") or "").strip(),
        "Assigned Recruiter": str(raw.get("Recruiter") or raw.get("TAQ Recruiter") or "").strip(),
        "Sourcer": str(raw.get("Sourcer") or "").strip(),
        "Taggd PM": str(raw.get("Taggd PM") or "").strip(),
        "BHR / HRBP": str(raw.get("Business HRBP") or "").strip(),
        "Hiring Agency": str(raw.get("Hiring Agency") or "").strip(),
        "TTO Days": raw.get("TTO") if not _is_empty(raw.get("TTO")) else None,
        "TTF Days": raw.get("TTF") if not _is_empty(raw.get("TTF")) else None,
        "Ageing Days": raw.get("Ageing") if not _is_empty(raw.get("Ageing")) else None,
        "FY Label": str(raw.get("Financial Year") or "").strip() or None,
        "Current Stage": str(raw.get("Stage") or "").strip() or None,
    }


def generate() -> Path:
    if not RAW_PM.exists():
        raise FileNotFoundError(f"Missing raw PM file: {RAW_PM}")
    shell = SHELL if SHELL.exists() else TEMPLATE_FALLBACK
    if not shell.exists():
        raise FileNotFoundError(f"Missing template shell: {shell}")

    raw_df = pd.read_excel(RAW_PM, sheet_name="Sheet1")
    deduped = _dedupe_raw(raw_df)
    req_ids = _assign_req_ids(deduped)

    rows: list[dict] = []
    for (_, raw), req_id in zip(deduped.iterrows(), req_ids):
        rows.append(_raw_row_to_template(raw, req_id))

    status_order = {"Joined": 0, "Offered": 1, "Open": 2, "On Hold": 3, "Cancelled": 4}
    rows.sort(
        key=lambda r: (
            status_order.get(str(r.get("Current Status") or ""), 9),
            re.sub(r"-\d+$", "", _norm_id(str(r.get("Req ID") or ""))),
        )
    )

    wb = openpyxl.load_workbook(shell)
    ws = wb["Position Tracker"]
    headers = _header_map(ws)
    _clear_data_area(ws, DATA_START, DATA_START + len(rows) + 10)

    for idx, row in enumerate(rows):
        excel_row = DATA_START + idx
        status = row.get("Current Status") or ""
        sjt = row.get("Source Joiner Type") or ""
        ws.cell(excel_row, 1).value = f"{status} · {sjt}"[:60]
        for key, col in headers.items():
            val = row.get(key)
            if val is not None and not (isinstance(val, float) and pd.isna(val)):
                if isinstance(val, str) and not val.strip():
                    continue
                ws.cell(excel_row, col).value = val

    _write_contractual(wb)
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    OUT_COPY.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUT_PATH, OUT_COPY)

    joined = [r for r in rows if r.get("Current Status") == "Joined"]
    joined_no_ctc = sum(1 for r in joined if _is_empty(r.get("Offered CTC (Lakhs)")))
    joined_no_jd = sum(1 for r in joined if _is_empty(r.get("Joining Date")))
    status_counts = Counter(str(r.get("Current Status")) for r in rows)
    sjt_counts = Counter(str(r.get("Source Joiner Type")) for r in rows)

    print(f"Raw PM rows: {len(raw_df)} → deduped: {len(deduped)} → wrote: {len(rows)}")
    print(f"Output → {OUT_PATH}")
    print(f"Copy → {OUT_COPY}")
    print(f"Status: {dict(status_counts)}")
    print(f"Source Joiner Type: {dict(sjt_counts)}")
    print(f"Joined: {len(joined)} | missing CTC: {joined_no_ctc} | missing join date: {joined_no_jd}")
    return OUT_PATH


if __name__ == "__main__":
    generate()
