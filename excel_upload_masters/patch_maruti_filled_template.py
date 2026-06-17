#!/usr/bin/env python3
"""
Patch project_1_tracker_template (1).xlsx using Maruti.xlsx (raw PM).

Reads HM baseline from project_1_tracker_template (2).xlsx (pipeline status preserved).
Fills dates, team, stage, CTC from Maruti.xlsx without changing Current Status.

Usage (from revagent/):
  PYTHONPATH=. python3 excel_upload_masters/patch_maruti_filled_template.py
"""
from __future__ import annotations

import datetime as dt
import re
import shutil
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
TAGGD_ROOT = ROOT.parent
FILLED = TAGGD_ROOT / "project_1_tracker_template (2).xlsx"
STATUS_REFERENCE = TAGGD_ROOT / "project_1_tracker_template (2).xlsx"
RAW_PM = TAGGD_ROOT / "Maruti.xlsx"
OUT_PATCHED = TAGGD_ROOT / "project_1_tracker_template (1).xlsx"
OUT_COPY = TAGGD_ROOT / "Maruti Suzuki Tracker - Patched.xlsx"
OUT_REPO = ROOT / "excel_files" / "Maruti Suzuki Tracker - Patched.xlsx"

DATA_START = 4
HIRE_TYPE_COL = "Type of Hire\n(Lateral, AIT, Campus, IJP)"

_RAW_STATUS_MAP = {
    "joined": "Joined",
    "ready to join": "Offered",
    "loi pending": "Offered",
    "salary fitment": "Offered",
    "documentation": "Offered",
    "sourcing": "Open",
    "intake": "Open",
    "interview": "Open",
    "cv sent": "Open",
    "final interview": "Open",
    "on_hold": "On Hold",
    "cancelled": "Cancelled",
}
_RAW_STAGE_MAP = {
    "joined": "Joined",
    "offered": "Offered",
    "on_hold": "On Hold",
    "cancelled": "Cancelled",
    "in_process": "Open",
    "selection": "Offered",
}
_STATUS_PRIORITY = {
    "joined": 6,
    "ready to join": 5,
    "documentation": 4,
    "salary fitment": 4,
    "loi pending": 4,
    "interview": 3,
    "cv sent": 3,
    "sourcing": 2,
    "intake": 2,
    "final interview": 3,
    "on_hold": 1,
    "cancelled": 0,
}


def _is_empty(val) -> bool:
    if val is None:
        return True
    if isinstance(val, float) and pd.isna(val):
        return True
    s = str(val).strip()
    return s in ("", "nan", "None", "NaT", "-", "—")


def _norm_id(val) -> str:
    return str(val).strip().lstrip("\xa0")


def _base_id(val) -> str:
    return re.sub(r"-\d+$", "", _norm_id(val))


def _norm_cand(val) -> str:
    s = str(val or "").strip().lower()
    if s.startswith("open —") or s.startswith("open -"):
        return ""
    return "" if s in ("", "-", "nan", "none") else s


def _to_date(val):
    if _is_empty(val):
        return None
    if isinstance(val, (dt.datetime, dt.date, pd.Timestamp)):
        return val.date() if hasattr(val, "date") else val
    s = str(val).strip()
    if s.replace(".", "", 1).isdigit():
        n = float(s)
        if n > 40000:
            return dt.date(1899, 12, 30) + dt.timedelta(days=int(n))
    try:
        parsed = pd.to_datetime(val, errors="coerce")
        if pd.notna(parsed):
            return parsed.date()
    except Exception:
        pass
    return None


def _is_serial_date(val) -> bool:
    if _is_empty(val):
        return False
    s = str(val).strip()
    return s.replace(".", "", 1).isdigit() and float(s) > 40000


def _ctc_lakhs(val) -> float | None:
    if _is_empty(val):
        return None
    try:
        v = float(str(val).replace(",", ""))
    except ValueError:
        return None
    if v <= 0:
        return None
    return round(v / 100_000, 2) if v > 500 else round(v, 2)


def _map_status(raw: pd.Series) -> str:
    st = str(raw.get("Stage") or "").strip().lower().replace(" ", "_")
    if st in _RAW_STAGE_MAP:
        return _RAW_STAGE_MAP[st]
    rs = str(raw.get("Status") or "").strip().lower().replace(" ", "_")
    return _RAW_STATUS_MAP.get(rs, "Open")


def _map_sjt(raw: pd.Series) -> str:
    soh = str(raw.get("Source of Hire") or "").strip().lower()
    ht = str(raw.get(HIRE_TYPE_COL) or "").strip().lower()
    if ht == "ijp" or soh == "ijp":
        return "IJP – Internal Job Posting"
    if soh == "er":
        return "ER – Employee Referral"
    if ht == "campus" or soh == "campus":
        return "Campus"
    if soh == "rehire" or "transfer" in ht:
        return "Internal Transfer"
    return "Taggd RPO"


def _status_priority(raw: pd.Series) -> int:
    rs = str(raw.get("Status") or "").strip().lower()
    st = str(raw.get("Stage") or "").strip().lower().replace(" ", "_")
    return max(_STATUS_PRIORITY.get(rs, 0), _STATUS_PRIORITY.get(st, 0))


def _header_map(ws) -> dict[str, int]:
    out: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        h = ws.cell(2, col).value
        if h and str(h) != "Scenario (example rows only)":
            out[str(h)] = col
    return out


def _set_cell(ws, headers: dict[str, int], row_num: int, key: str, val) -> None:
    col = headers.get(key)
    if col is None or val is None:
        return
    if isinstance(val, str) and not val.strip():
        return
    if isinstance(val, float) and pd.isna(val):
        return
    cell = ws.cell(row_num, col)
    if type(cell).__name__ == "MergedCell":
        return
    cell.value = val


def _req_suffix(req_id) -> int:
    m = re.search(r"-(\d+)$", _norm_id(req_id))
    return int(m.group(1)) if m else 0


def _build_raw_by_job(raw_df: pd.DataFrame) -> dict[str, list[tuple[int, pd.Series]]]:
    by_job: dict[str, list[tuple[int, pd.Series]]] = defaultdict(list)
    for idx, raw in raw_df.iterrows():
        jid = _norm_id(raw.get("Job ID"))
        if jid:
            by_job[jid].append((idx, raw))
    return by_job


def _build_row_matches(df: pd.DataFrame, raw_df: pd.DataFrame) -> list[pd.Series | None]:
    """
    Match each template row to one raw PM row.
    1) Named candidates match by (Job ID, Candidate Name).
    2) Open / placeholder rows match remaining raw lines in suffix order.
    """
    by_job = _build_raw_by_job(raw_df)
    used: dict[str, set[int]] = defaultdict(set)
    n = len(df)
    matches: list[pd.Series | None] = [None] * n

    # Pass 1 — named candidates
    for pos in range(n):
        row = df.iloc[pos]
        bid = _base_id(row["Req ID"])
        cand = _norm_cand(row["Candidate Name"])
        if not cand:
            continue
        for idx, raw in by_job.get(bid, []):
            if idx in used[bid]:
                continue
            if _norm_cand(raw.get("Candidate Name")) == cand:
                matches[pos] = raw
                used[bid].add(idx)
                break

    # Pass 2 — positional for open / unmatched rows
    for pos in range(n):
        if matches[pos] is not None:
            continue
        row = df.iloc[pos]
        bid = _base_id(row["Req ID"])
        pool = [(idx, raw) for idx, raw in by_job.get(bid, []) if idx not in used[bid]]
        suffix = _req_suffix(row["Req ID"])
        pick = pool[suffix] if suffix < len(pool) else (pool[0] if pool else None)
        if pick:
            used[bid].add(pick[0])
            matches[pos] = pick[1]

    return matches


_STAGE_TO_PIPELINE = {
    "In_Process": "Open",
    "Joined": "Joined",
    "Offered": "Offered",
    "On_Hold": "On Hold",
    "Cancelled": "Cancelled",
    "Selection": "Offered",
}


def _pipeline_status(raw: pd.Series) -> str:
    """Pipeline status (Current Status) from raw Stage."""
    st = str(raw.get("Stage") or "").strip()
    return _STAGE_TO_PIPELINE.get(st, "Open")


def _patch_row(row: dict, raw: pd.Series | None) -> tuple[dict, Counter]:
    changes: Counter = Counter()
    if raw is None:
        return row, changes

    out = dict(row)
    status = str(out.get("Current Status") or "").strip()
    req_id = _norm_id(out.get("Req ID"))

    # --- Pipeline status: never overwrite (set from STATUS_REFERENCE in patch()) ---

    # --- Joining Date: fill only when empty in HM (never overwrite existing values) ---
    raw_join = _to_date(raw.get("Joining Date"))
    if raw_join and _is_empty(out.get("Joining Date")) and status in ("Joined", "Offered"):
        out["Joining Date"] = raw_join
        changes["joining_date"] += 1

    if status == "Joined" and _is_empty(out.get("Closure Date")) and raw_join:
        out["Closure Date"] = raw_join
        changes["closure_date"] += 1

    # --- CTC: fill empty on Joined/Offered ---
    raw_ctc = _ctc_lakhs(raw.get("Offered CTC"))
    if raw_ctc and status in ("Joined", "Offered") and _is_empty(out.get("Offered CTC (Lakhs)")):
        out["Offered CTC (Lakhs)"] = raw_ctc
        changes["ctc_filled"] += 1

    # --- Candidate Name on Open ---
    if status == "Open" and _is_empty(out.get("Candidate Name")):
        cand = str(raw.get("Candidate Name") or "").strip()
        if cand and cand != "-":
            out["Candidate Name"] = cand
            changes["candidate_filled"] += 1
        else:
            out["Candidate Name"] = f"Open — {req_id}"
            changes["candidate_placeholder"] += 1

    # --- Fill-if-empty columns ---
    fill_map = [
        ("Intake Date", "Intake Date", _to_date),
        ("Candidate Selection Date", "Selection Date", _to_date),
        ("Offered Date", "LOI Date", _to_date),
        ("Offer Accepted Date", "LOI Date", _to_date),
        ("LOI Date", "LOI Date", _to_date),
        ("Source of Hire", "Source of Hire", lambda v: str(v).strip() if not _is_empty(v) else None),
        ("Sourcer", "Sourcer", lambda v: str(v).strip() if not _is_empty(v) else None),
        ("Taggd PM", "Taggd PM", lambda v: str(v).strip() if not _is_empty(v) else None),
        ("BHR / HRBP", "Business HRBP", lambda v: str(v).strip() if not _is_empty(v) else None),
        ("Hiring Agency", "Hiring Agency", lambda v: str(v).strip() if not _is_empty(v) else None),
        ("TTO Days", "TTO", lambda v: v if not _is_empty(v) else None),
        ("TTF Days", "TTF", lambda v: v if not _is_empty(v) else None),
        ("Ageing Days", "Ageing", lambda v: v if not _is_empty(v) else None),
        ("FY Label", "Financial Year", lambda v: str(v).strip() if not _is_empty(v) else None),
        ("Current Stage", "Stage", lambda v: str(v).strip() if not _is_empty(v) else None),
    ]
    for tcol, rcol, fn in fill_map:
        if _is_empty(out.get(tcol)):
            val = fn(raw.get(rcol))
            if val is not None and not (isinstance(val, str) and not str(val).strip()):
                out[tcol] = val
                changes[f"filled_{tcol}"] += 1

    # Source Joiner Type — refresh from raw if empty (usually already set)
    if _is_empty(out.get("Source Joiner Type")):
        out["Source Joiner Type"] = _map_sjt(raw)
        changes["source_joiner_type"] += 1

    return out, changes


def _preserve_hm_row(patched: dict, ref: dict) -> dict:
    """Restore every non-empty value from the HM baseline row (2).xlsx."""
    out = dict(patched)
    for col, val in ref.items():
        if col == "Scenario (example rows only)":
            continue
        if str(col).startswith("Unnamed"):
            continue
        if not _is_empty(val):
            out[col] = val
    return out


def patch() -> Path:
    if not FILLED.exists():
        raise FileNotFoundError(f"Missing filled template: {FILLED}")
    if not RAW_PM.exists():
        raise FileNotFoundError(f"Missing raw PM: {RAW_PM}")

    import sys

    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from backend.core.tracker_sheet_io import read_tracker_sheet_dataframe

    df, _ = read_tracker_sheet_dataframe(str(FILLED), "Position Tracker")
    original_count = len(df)

    # Preserve all HM-filled cells from reference file (project_1_tracker_template (2).xlsx)
    ref_by_req: dict[str, dict] = {}
    if STATUS_REFERENCE.exists():
        ref_df, _ = read_tracker_sheet_dataframe(str(STATUS_REFERENCE), "Position Tracker")
        for _, ref_row in ref_df.iterrows():
            rid = _norm_id(ref_row.get("Req ID"))
            if rid:
                ref_by_req[rid] = ref_row.to_dict()

    raw_df = pd.read_excel(RAW_PM, sheet_name="Sheet1")
    row_matches = _build_row_matches(df, raw_df)

    total_changes: Counter = Counter()
    patched_rows: list[dict] = []
    for pos in range(len(df)):
        row = df.iloc[pos]
        patched, ch = _patch_row(row.to_dict(), row_matches[pos])
        rid = _norm_id(patched.get("Req ID"))
        if rid in ref_by_req:
            patched = _preserve_hm_row(patched, ref_by_req[rid])
            ch["hm_preserved"] += 1
        patched_rows.append(patched)
        total_changes.update(ch)

    assert len(patched_rows) == original_count == 579, f"Row count changed: {len(patched_rows)}"

    wb = openpyxl.load_workbook(FILLED)
    ws = wb["Position Tracker"]
    headers = _header_map(ws)

    for idx, row in enumerate(patched_rows):
        excel_row = DATA_START + idx
        status = row.get("Current Status") or ""
        sjt = row.get("Source Joiner Type") or ""
        banner = ws.cell(excel_row, 1)
        if type(banner).__name__ != "MergedCell":
            banner.value = f"{status} · {sjt}"[:60]
        for key, col in headers.items():
            val = row.get(key)
            if val is None or (isinstance(val, float) and pd.isna(val)):
                continue
            if isinstance(val, str) and not val.strip():
                continue
            _set_cell(ws, headers, excel_row, key, val)

    wb.save(OUT_PATCHED)
    shutil.copy2(OUT_PATCHED, OUT_COPY)
    OUT_REPO.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUT_PATCHED, OUT_REPO)

    # Validation summary
    joined = [r for r in patched_rows if r.get("Current Status") == "Joined"]
    serial = sum(1 for r in joined if _is_serial_date(r.get("Joining Date")))
    no_ctc = sum(1 for r in joined if _is_empty(r.get("Offered CTC (Lakhs)")))
    rev_ready = sum(
        1
        for r in joined
        if not _is_serial_date(r.get("Joining Date"))
        and not _is_empty(r.get("Joining Date"))
        and not _is_empty(r.get("Offered CTC (Lakhs)"))
    )

    print(f"Patched {original_count} rows → {OUT_PATCHED}")
    print(f"Also → {OUT_COPY}")
    print(f"Copy → {OUT_REPO}")
    print("\nChanges applied:")
    for k, v in total_changes.most_common():
        print(f"  {k}: {v}")
    status_counts = Counter(str(r.get("Current Status") or "") for r in patched_rows)
    stage_counts = Counter(str(r.get("Current Stage") or "") for r in patched_rows)
    print(f"\nPipeline status (Current Status) — use this for pivot:")
    for label in ("Open", "Joined", "Offered", "On Hold", "Cancelled"):
        print(f"  {label}: {status_counts.get(label, 0)}")
    print(f"  Grand Total: {sum(status_counts.values())}")
    print(f"\nRaw Stage breakdown:")
    for label, cnt in stage_counts.most_common():
        print(f"  {label}: {cnt}")
    print(f"\nJoined: {len(joined)} | revenue-ready: {rev_ready}/{len(joined)}")
    print(f"Serial join dates remaining: {serial}")
    print(f"Joined missing CTC: {no_ctc} (not in raw PM — manual entry needed)")
    return OUT_PATCHED


if __name__ == "__main__":
    patch()
