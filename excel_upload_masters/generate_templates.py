#!/usr/bin/env python3
"""
Generate blank Excel upload masters from SQLAlchemy models (column names = DB field names).

Run from repo root:
  python excel_upload_masters/generate_templates.py

Requires: openpyxl (from project requirements.txt)
"""
from __future__ import annotations

import os
import sys
from datetime import datetime

# Repo root on path
ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, ROOT)

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from backend.db import database as db

_EXCEL_DIR = os.path.dirname(os.path.abspath(__file__))
if _EXCEL_DIR not in sys.path:
    sys.path.insert(0, _EXCEL_DIR)
from column_dropdowns import get_dropdown_options

OUT_DIR = os.path.join(os.path.dirname(__file__), "templates")

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
NOTE_FONT = Font(italic=True, size=9, color="6B7280")

# (file_stem, [(sheet_title, Model), ...]) — order reflects recommended load order for FKs
WORKBOOKS: list[tuple[str, list[tuple[str, type]]]] = [
    (
        "01_spine_clients_projects",
        [
            ("clients", db.Client),
            ("projects", db.Project),
        ],
    ),
    (
        "02_users_rbac",
        [
            ("users", db.User),
            ("user_project_assignments", db.UserProjectAssignment),
        ],
    ),
    (
        "03_commercial_contracts",
        [
            ("project_contracts", db.ProjectContract),
        ],
    ),
    (
        "04_client_onboarding_transitions",
        [
            ("project_transitions", db.ProjectTransition),
        ],
    ),
    (
        "05_pipeline_requisitions_records",
        [
            ("records", db.Record),
        ],
    ),
    (
        "06_candidates",
        [
            ("candidates", db.Candidate),
        ],
    ),
    (
        "07_candidate_identity_masters",
        [
            ("candidate_masters", db.CandidateMaster),
            ("candidate_master_links", db.CandidateMasterLink),
        ],
    ),
    (
        "08_sla",
        [
            ("metric_definitions", db.MetricDefinition),
            ("sla_performances", db.SLAPerformance),
        ],
    ),
    (
        "09_workforce_management",
        [
            ("wfm_hr_benchmarks", db.WFMHRBenchmark),
            ("wfm_resource_gaps", db.WFMResourceGap),
        ],
    ),
    (
        "10_finance_core",
        [
            ("finance_monthly_ledger", db.FinanceMonthlyLedger),
            ("finance_cash_flow", db.FinanceCashFlow),
            ("finance_efficiency_kpis", db.FinanceEfficiencyKPI),
        ],
    ),
    (
        "11_revenue_trackers",
        [
            ("revenue_weekly_submission", db.RevenueWeeklySubmission),
            ("revenue_forecast_weekly", db.RevenueForecastWeekly),
            ("revenue_visibility_snapshot", db.RevenueVisibilitySnapshot),
        ],
    ),
    (
        "12_billing_taggd_workflow",
        [
            ("taggd_revenue_billing", db.TaggdRevenueBilling),
            ("finance_billing_workflow", db.FinanceBillingWorkflow),
            ("finance_billing_validation_events", db.FinanceBillingValidationEvent),
            ("finance_payment_receipts", db.FinancePaymentReceipt),
            ("finance_tds_certificates", db.FinanceTdsCertificate),
            ("finance_bank_statement_lines", db.FinanceBankStatementLine),
        ],
    ),
    (
        "13_meetings",
        [
            ("platform_meetings", db.Meeting),
            ("meeting_action_items", db.MeetingActionItem),
        ],
    ),
    (
        "14_tasks",
        [
            ("platform_tasks", db.Task),
            ("task_assignees", db.TaskAssignee),
        ],
    ),
    (
        "15_vendor_resume_licenses",
        [
            ("resume_supplier_licenses", db.ResumeSupplierLicense),
        ],
    ),
    (
        "16_ingestion_audit",
        [
            ("ingestion_events", db.IngestionEvent),
        ],
    ),
]

SENSITIVE_COLUMNS = {
    "password_hash": "Do not store plaintext. Leave blank for invite-only, or use bcrypt hash if importing via tool.",
}

DATA_START_ROW = 3
MAX_DATA_ROWS = 5000
LIST_CACHE_MAX = 200  # avoid Excel / performance issues for huge inline ranges


class _ListAllocator:
    """Hidden `lists` sheet: one column per *unique* option list, reused when identical."""

    def __init__(self, wb: Workbook) -> None:
        self._wb = wb
        if "lists" in wb.sheetnames:
            self._ws = wb["lists"]
        else:
            self._ws = wb.create_sheet("lists")
            self._ws.sheet_state = "hidden"
        self._col = 0
        self._cache: dict[tuple[str, ...], str] = {}

    def ref_for(self, options: list[str]) -> str | None:
        if not options or len(options) > LIST_CACHE_MAX:
            return None
        key = tuple(str(x) if x is not None else "" for x in options)
        if key in self._cache:
            return self._cache[key]
        self._col += 1
        c = self._col
        for i, v in enumerate(options, start=1):
            val = v if (v is not None and v != "") else None
            self._ws.cell(row=i, column=c, value=val)
        lo = get_column_letter(c)
        n = len(options)
        ref = f"lists!${lo}$1:${lo}${n}"
        self._cache[key] = ref
        return ref


def _columns_for_model(model) -> list[str]:
    # Table column order matches physical / migration order in SQLite.
    return [c.key for c in model.__table__.columns]


def _build_sheet(wb: Workbook, model, sheet_name: str, list_alloc: _ListAllocator) -> None:
    ws = wb.create_sheet(title=sheet_name[:31])  # Excel limit
    table = model.__table__
    cols = _columns_for_model(model)
    for j, col in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=j, value=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        note = SENSITIVE_COLUMNS.get(col, "")
        if "json" in col.lower() or col.endswith("_json"):
            note = (note + " | JSON (object/array as text). ").strip(" |")
        col_t = table.c[col].type
        col_type = str(col_t)
        if col.endswith("_at") or col.endswith("_date") or "DATETIME" in col_type or "DateTime" in col_type:
            if not note:
                note = "Date/datetime — ISO-8601 or Excel date"
        if "JSON" in col_type and "json" not in col.lower():
            note = (note + " | JSON. ").strip(" |")
        opts = get_dropdown_options(model, col, col_t)
        if opts:
            note = (note + " | Pick from list (dropdown)").strip(" |")
        ws.cell(row=2, column=j, value=note)
        ws.cell(row=2, column=j).font = NOTE_FONT
    ws.freeze_panes = f"A{DATA_START_ROW}"
    for j, col in enumerate(cols, start=1):
        col_t = table.c[col].type
        opts = get_dropdown_options(model, col, col_t)
        if not opts:
            continue
        ref = list_alloc.ref_for(opts)
        if not ref:
            continue
        lo = get_column_letter(j)
        cr = f"{lo}{DATA_START_ROW}:{lo}{DATA_START_ROW + MAX_DATA_ROWS - 1}"
        dv = DataValidation(
            type="list",
            formula1=f"={ref}",
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Invalid value",
            error="Choose a value from the list, or clear the cell.",
        )
        ws.add_data_validation(dv)
        dv.add(cr)
    # Widen a bit
    for j in range(1, min(len(cols) + 1, 30)):
        ws.column_dimensions[get_column_letter(j)].width = 18


def main() -> None:
    os.makedirs(OUT_DIR, exist_ok=True)
    for stem, sheets in WORKBOOKS:
        wb = Workbook()
        # remove default sheet
        default = wb.active
        wb.remove(default)
        list_alloc = _ListAllocator(wb)
        for sheet_label, model in sheets:
            safe_name = sheet_label[:31]
            _build_sheet(wb, model, safe_name, list_alloc)
        if not wb.sheetnames:
            continue
        path = os.path.join(OUT_DIR, f"{stem}.xlsx")
        wb.save(path)
        print("Wrote", path)

    # Index manifest
    manifest = os.path.join(OUT_DIR, "_generated_at.txt")
    with open(manifest, "w", encoding="utf-8") as f:
        f.write(f"Generated: {datetime.utcnow().isoformat()}Z\n")
        f.write("Source: SQLAlchemy models in backend/db/database.py\n")
    print("Done.")


if __name__ == "__main__":
    main()
